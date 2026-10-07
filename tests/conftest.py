import pandas as pd
import pytest
from openpyxl import Workbook
from openpyxl.worksheet.table import Table


@pytest.fixture
def sample_xlsx(tmp_path):
    """Workbook with sheet 'Sales' (holding Table 'SalesTable'), 'Notes', and an empty 'Blank'."""
    wb = Workbook()
    sales = wb.active
    sales.title = "Sales"
    sales.append(["id", "amount"])
    sales.append([1, 10.5])
    sales.append([2, 20.0])
    sales.append([3, None])
    sales.add_table(Table(displayName="SalesTable", ref="A1:B4"))
    notes = wb.create_sheet("Notes")
    notes.append(["note"])
    notes.append(["hello"])
    wb.create_sheet("Blank")
    path = tmp_path / "book.xlsx"
    wb.save(path)
    return path


@pytest.fixture
def sample_ods(tmp_path):
    path = tmp_path / "book.ods"
    with pd.ExcelWriter(path, engine="odf") as writer:
        pd.DataFrame({"a": [1, 2], "b": ["x", "y"]}).to_excel(writer, sheet_name="First", index=False)
        pd.DataFrame({"c": [3]}).to_excel(writer, sheet_name="Second", index=False)
    return path
