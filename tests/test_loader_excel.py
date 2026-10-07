from pathlib import Path

import pytest

from qxc.loader import LoaderError, list_sheets, load_table

SAMPLE_XLS = Path(__file__).parent / "data" / "sample.xls"


def test_xlsx_sheets_sizes_and_tables(sample_xlsx):
    sheets = {s.name: s for s in list_sheets(sample_xlsx)}
    assert list(sheets) == ["Sales", "Notes", "Blank"]
    assert (sheets["Sales"].rows, sheets["Sales"].cols) == (4, 2)
    assert sheets["Sales"].tables == ("SalesTable",)
    assert sheets["Notes"].tables == ()
    assert (sheets["Blank"].rows, sheets["Blank"].cols) == (0, 0)


def test_xlsx_load_sheet_and_default_sheet(sample_xlsx):
    assert load_table(sample_xlsx, sheet="Notes").shape == (1, 1)
    assert load_table(sample_xlsx).shape == (3, 2)  # first sheet


def test_xlsx_load_named_table(sample_xlsx):
    df = load_table(sample_xlsx, sheet="Sales", table="SalesTable")
    assert df.shape == (3, 2)
    assert list(df.columns) == ["id", "amount"]


def test_xlsx_missing_table_is_a_loader_error(sample_xlsx):
    with pytest.raises(LoaderError):
        load_table(sample_xlsx, sheet="Sales", table="Nope")


def test_ods_sheets_and_load(sample_ods):
    assert [s.name for s in list_sheets(sample_ods)] == ["First", "Second"]
    assert list_sheets(sample_ods)[0].tables == ()
    assert load_table(sample_ods, sheet="First").shape == (2, 2)


def test_password_protected_xlsx_message(tmp_path):
    path = tmp_path / "locked.xlsx"
    path.write_bytes(b"\xd0\xcf\x11\xe0" + b"\x00" * 64)
    with pytest.raises(LoaderError, match="password"):
        list_sheets(path)


def test_corrupt_xlsx_is_a_loader_error(tmp_path):
    path = tmp_path / "bad.xlsx"
    path.write_bytes(b"this is not a zip file")
    with pytest.raises(LoaderError):
        list_sheets(path)
    with pytest.raises(LoaderError):
        load_table(path)


@pytest.mark.skipif(not SAMPLE_XLS.exists(), reason="Jose has not supplied tests/data/sample.xls")
def test_xls_reads():
    sheets = list_sheets(SAMPLE_XLS)
    assert sheets
    assert not load_table(SAMPLE_XLS, sheet=sheets[0].name).empty


def test_list_sheets_does_not_load_every_cell(sample_xlsx, monkeypatch):  # review #3
    import openpyxl

    def boom(*args, **kwargs):
        raise AssertionError("list_sheets must not load the whole workbook")

    monkeypatch.setattr(openpyxl, "load_workbook", boom)
    sheets = {s.name: s for s in list_sheets(sample_xlsx)}
    assert (sheets["Sales"].rows, sheets["Sales"].cols) == (4, 2)
    assert sheets["Sales"].tables == ("SalesTable",)
    assert (sheets["Blank"].rows, sheets["Blank"].cols) == (0, 0)


def test_xlsx_sheet_with_one_cell_and_special_characters(tmp_path):  # review #3
    from openpyxl import Workbook
    from openpyxl.worksheet.table import Table

    wb = Workbook()
    ws = wb.active
    ws.title = "R&D <1>"
    ws.append(["k", "v"])
    ws.append(["a", 1])
    ws.add_table(Table(displayName="R_D", ref="A1:B2"))
    one = wb.create_sheet("One")
    one["A1"] = "x"
    path = tmp_path / "odd.xlsx"
    wb.save(path)
    sheets = {s.name: s for s in list_sheets(path)}
    assert sheets["R&D <1>"].tables == ("R_D",)
    assert (sheets["One"].rows, sheets["One"].cols) == (1, 1)
    assert load_table(path, sheet="R&D <1>", table="R_D").shape == (1, 2)
