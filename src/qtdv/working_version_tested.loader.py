"""Read table files with pandas.

This is the only module that opens data files. It hides the differences
between CSV, Excel and LibreOffice Calc files from the rest of the app.
Phase 1 is read-only.
"""

from __future__ import annotations

import csv
import posixpath
import re
import warnings
import zipfile
from dataclasses import dataclass
from pathlib import Path
from xml.etree import ElementTree as ET

import pandas as pd

LARGE_FILE_BYTES = 50 * 1024 * 1024
_LEADING_ZERO = re.compile(r"^0\d")
_OLE_MAGIC = b"\xd0\xcf\x11\xe0"  # an .xlsx that starts like this is encrypted
_XLSX_EXTS = (".xlsx", ".xlsm")
_LONG_DIGITS = re.compile(r"\d{16,}")  # too long for a float to hold exactly
_NS = {
    "m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}


class LoaderError(Exception):
    """A problem reading a file. The message is written for the user to read."""


@dataclass(frozen=True)
class SheetInfo:
    """A sheet in a workbook. ``rows``/``cols`` are the used range, when known."""

    name: str
    rows: int | None = None
    cols: int | None = None
    tables: tuple[str, ...] = ()


def is_large(size_bytes: int) -> bool:
    """True when a file is over the 50 MB warning limit."""
    return size_bytes > LARGE_FILE_BYTES


def clean_columns(df: pd.DataFrame) -> pd.DataFrame:
    """Make every column name a unique, non-blank string."""
    used: set[str] = set()
    names: list[str] = []
    for position, col in enumerate(df.columns):
        blank = col is None or (isinstance(col, float) and pd.isna(col)) or not str(col).strip()
        base = f"Unnamed: {position}" if blank else str(col).strip()
        name, n = base, 0
        while name in used:
            n += 1
            name = f"{base}.{n}"
        used.add(name)
        names.append(name)
    df.columns = names
    return df


def _infer_types(df: pd.DataFrame) -> pd.DataFrame:
    """Turn all-text CSV columns into numbers when that is safe.

    A column stays text if any value starts with 0 followed by a digit
    (``00123``), so ZIP codes and case numbers keep their leading zeros.
    """
    for col in df.columns:
        values = df[col].dropna()
        if values.empty or values.str.match(_LEADING_ZERO).any():
            continue
        if values.str.fullmatch(_LONG_DIGITS).any():
            continue
        converted = pd.to_numeric(df[col], errors="coerce")
        if converted.notna().sum() == df[col].notna().sum():
            df[col] = converted
    return df


def _read_csv(path: Path) -> pd.DataFrame:
    last_error: Exception | None = None
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            with warnings.catch_warnings():
                # Extra fields on a row are dropped (index_col=False); pandas warns about that,
                # and a warning printed over the terminal screen would be noise.
                warnings.simplefilter("ignore", pd.errors.ParserWarning)
                try:
                    return _infer_types(pd.read_csv(path, dtype=str, encoding=encoding, index_col=False))
                except pd.errors.ParserError:
                    # A row further down has more fields than the header (for example an unquoted comma
                    # inside a note). The fast reader gives up; the slower one can cut the extra fields.
                    return _read_csv_cutting_extra_fields(path, encoding)
        except UnicodeDecodeError as exc:
            last_error = exc
        except pd.errors.EmptyDataError:
            return pd.DataFrame()
        except pd.errors.ParserError as exc:
            raise LoaderError(f"Cannot parse CSV: {exc}") from exc
    raise LoaderError(f"Unknown text encoding: {last_error}")


def _read_csv_cutting_extra_fields(path: Path, encoding: str) -> pd.DataFrame:
    """Read a CSV whose rows are longer than its header, cutting each row to the header's width.

    The result carries a note (``df.attrs["notes"]``) so the app can tell the user what was cut.
    """
    df = _infer_types(pd.read_csv(path, dtype=str, encoding=encoding, index_col=False, engine="python"))
    width = len(df.columns)
    with open(path, newline="", encoding=encoding) as handle:
        reader = csv.reader(handle)
        long_rows = [reader.line_num for row in reader if len(row) > width]
    plural = "s" if len(long_rows) != 1 else ""
    first = f" (first at line {long_rows[0]})" if long_rows else ""
    df.attrs["notes"] = [f"{len(long_rows)} row{plural} had extra fields and the extra values were cut{first}"]
    return df


def _wrap(path: Path, exc: Exception) -> LoaderError:
    """Turn a low-level exception into a message a person can act on."""
    if isinstance(exc, PermissionError):
        return LoaderError("File in use, try again")
    if "encrypted" in str(exc).lower() or (path.suffix.lower() in _XLSX_EXTS and _looks_encrypted(path)):
        return LoaderError("Workbook is password-protected and cannot be read")
    return LoaderError(f"Cannot read {path.name}: {type(exc).__name__}: {exc}")


def _looks_encrypted(path: Path) -> bool:
    try:
        with open(path, "rb") as handle:
            return handle.read(4) == _OLE_MAGIC
    except OSError:
        return False


def list_sheets(path: Path | str) -> list[SheetInfo]:
    """Sheets (and, for .xlsx, named Tables) in a file. CSV files have none."""
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".csv":
        return []
    try:
        if ext in _XLSX_EXTS:
            return _xlsx_sheets(path)
        if ext == ".xls":
            return _xls_sheets(path)
        if ext == ".ods":
            return _ods_sheets(path)
    except LoaderError:
        raise
    except Exception as exc:  # reader libraries raise many different types
        raise _wrap(path, exc) from exc
    raise LoaderError(f"Unsupported file type: {ext or path.name}")


def load_table(path: Path | str, sheet: str | None = None, table: str | None = None) -> pd.DataFrame:
    """Read a whole table into a DataFrame.

    ``sheet`` picks a sheet (first sheet when omitted). ``table`` picks a named
    Excel Table on that sheet (.xlsx only).
    """
    path = Path(path)
    ext = path.suffix.lower()
    try:
        if ext == ".csv":
            df = _read_csv(path)
        elif table is not None and ext in _XLSX_EXTS:
            df = _read_xlsx_table(path, sheet, table)
        elif ext in _XLSX_EXTS:
            df = pd.read_excel(path, sheet_name=sheet or 0, engine="openpyxl")
        elif ext == ".xls":
            df = pd.read_excel(path, sheet_name=sheet or 0, engine="xlrd")
        elif ext == ".ods":
            df = pd.read_excel(path, sheet_name=sheet or 0, engine="odf")
        else:
            raise LoaderError(f"Unsupported file type: {ext or path.name}")
    except LoaderError:
        raise
    except Exception as exc:  # reader libraries raise many different types
        raise _wrap(path, exc) from exc
    return clean_columns(df)


def _resolve(base_dir: str, target: str) -> str:
    """Turn a relationship target into a path inside the zip file."""
    if target.startswith("/"):
        return target[1:]
    return posixpath.normpath(posixpath.join(base_dir, target))


def _xlsx_layout(path: Path) -> list[tuple[str, str | None, dict[str, str]]]:
    """Read sheet names, used range and named Tables straight from the zip.

    No cell data is loaded, so this is fast even for large workbooks. Returns
    ``(sheet name, used-range text or None, {table name: table range})`` per sheet.
    """
    from openpyxl.utils.cell import range_boundaries  # noqa: F401  (validates the import early)

    layout = []
    with zipfile.ZipFile(path) as archive:
        workbook = ET.fromstring(archive.read("xl/workbook.xml"))
        rels = {r.get("Id"): r.get("Target") for r in ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))}
        for sheet in workbook.find("m:sheets", _NS):
            sheet_path = _resolve("xl", rels[sheet.get(f"{{{_NS['r']}}}id")])
            head = archive.read(sheet_path) if archive.getinfo(sheet_path).file_size <= 65536 else None
            if head is None:
                with archive.open(sheet_path) as handle:
                    head = handle.read(4096)
            match = re.search(rb'<dimension ref="([^"]+)"', head)
            ref = match.group(1).decode() if match else None
            if ref and re.fullmatch(r"([A-Z]+\d+)(:\1)?", ref) and not re.search(rb"<c[ >]", head):
                ref = ""  # a single-cell range with no cell inside: an empty sheet
            tables: dict[str, str] = {}
            rels_path = posixpath.join(posixpath.dirname(sheet_path), "_rels", posixpath.basename(sheet_path) + ".rels")
            if rels_path in archive.namelist():
                for rel in ET.fromstring(archive.read(rels_path)):
                    if rel.get("Type", "").endswith("/table"):
                        table = ET.fromstring(archive.read(_resolve(posixpath.dirname(sheet_path), rel.get("Target"))))
                        tables[table.get("displayName") or table.get("name")] = table.get("ref")
            layout.append((sheet.get("name"), ref, tables))
    return layout


def _xlsx_sheets(path: Path) -> list[SheetInfo]:
    from openpyxl.utils.cell import range_boundaries

    sheets = []
    for name, ref, tables in _xlsx_layout(path):
        if ref is None:
            rows = cols = None
        elif ref == "":
            rows = cols = 0
        else:
            _, _, max_col, max_row = range_boundaries(ref)
            rows, cols = max_row, max_col
        sheets.append(SheetInfo(name, rows, cols, tuple(tables)))
    return sheets


def _xls_sheets(path: Path) -> list[SheetInfo]:
    import xlrd

    book = xlrd.open_workbook(path, on_demand=True)
    try:
        return [SheetInfo(s.name, s.nrows, s.ncols) for s in (book.sheet_by_index(i) for i in range(book.nsheets))]
    finally:
        book.release_resources()


def _ods_sheets(path: Path) -> list[SheetInfo]:
    # Sizes are not available without reading each sheet, so only names are listed.
    with pd.ExcelFile(path, engine="odf") as workbook:
        return [SheetInfo(name) for name in workbook.sheet_names]


def _read_xlsx_table(path: Path, sheet: str | None, table: str) -> pd.DataFrame:
    from openpyxl import load_workbook
    from openpyxl.utils.cell import range_boundaries

    layout = _xlsx_layout(path)
    sheet_name = sheet or layout[0][0]
    ref = next((tables[table] for name, _, tables in layout if name == sheet_name and table in tables), None)
    if ref is None:
        raise LoaderError(f"Table '{table}' not found on sheet '{sheet_name}'")
    min_col, min_row, max_col, max_row = range_boundaries(ref)
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        rows = list(
            workbook[sheet_name].iter_rows(
                min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col, values_only=True
            )
        )
    finally:
        workbook.close()
    if not rows:
        return pd.DataFrame()
    return pd.DataFrame(rows[1:], columns=list(rows[0])).infer_objects()
