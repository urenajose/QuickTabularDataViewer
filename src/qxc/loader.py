"""Read table files with pandas.

This is the only module that opens data files. It hides the differences
between CSV, Excel and LibreOffice Calc files from the rest of the app.
Phase 1 is read-only.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

import pandas as pd

LARGE_FILE_BYTES = 50 * 1024 * 1024
_LEADING_ZERO = re.compile(r"^0\d")
_OLE_MAGIC = b"\xd0\xcf\x11\xe0"  # an .xlsx that starts like this is encrypted
_XLSX_EXTS = (".xlsx", ".xlsm")


class LoaderError(Exception):
    """A problem reading a file. The message is written for the user to read."""


@dataclass(frozen=True)
class SheetInfo:
    """A sheet in a workbook. ``rows``/``cols`` are the used range, when known."""

    name: str
    rows: int | None = None
    cols: int | None = None
    tables: tuple[str, ...] = ()


def is_large(size_bytes: int) -> bool:
    """True when a file is over the 50 MB warning limit."""
    return size_bytes > LARGE_FILE_BYTES


def clean_columns(df: pd.DataFrame) -> pd.DataFrame:
    """Make every column name a unique, non-blank string."""
    used: set[str] = set()
    names: list[str] = []
    for position, col in enumerate(df.columns):
        blank = col is None or (isinstance(col, float) and pd.isna(col)) or not str(col).strip()
        base = f"Unnamed: {position}" if blank else str(col).strip()
        name, n = base, 0
        while name in used:
            n += 1
            name = f"{base}.{n}"
        used.add(name)
        names.append(name)
    df.columns = names
    return df


def _infer_types(df: pd.DataFrame) -> pd.DataFrame:
    """Turn all-text CSV columns into numbers when that is safe.

    A column stays text if any value starts with 0 followed by a digit
    (``00123``), so ZIP codes and case numbers keep their leading zeros.
    """
    for col in df.columns:
        values = df[col].dropna()
        if values.empty or values.str.match(_LEADING_ZERO).any():
            continue
        converted = pd.to_numeric(df[col], errors="coerce")
        if converted.notna().sum() == df[col].notna().sum():
            df[col] = converted
    return df


def _read_csv(path: Path) -> pd.DataFrame:
    last_error: Exception | None = None
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return _infer_types(pd.read_csv(path, dtype=str, encoding=encoding))
        except UnicodeDecodeError as exc:
            last_error = exc
        except pd.errors.EmptyDataError:
            return pd.DataFrame()
        except pd.errors.ParserError as exc:
            raise LoaderError(f"Cannot parse CSV: {exc}") from exc
    raise LoaderError(f"Unknown text encoding: {last_error}")


def _wrap(path: Path, exc: Exception) -> LoaderError:
    """Turn a low-level exception into a message a person can act on."""
    if isinstance(exc, PermissionError):
        return LoaderError("File in use, try again")
    if "encrypted" in str(exc).lower() or (path.suffix.lower() in _XLSX_EXTS and _looks_encrypted(path)):
        return LoaderError("Workbook is password-protected and cannot be read")
    return LoaderError(f"Cannot read {path.name}: {type(exc).__name__}: {exc}")


def _looks_encrypted(path: Path) -> bool:
    try:
        with open(path, "rb") as handle:
            return handle.read(4) == _OLE_MAGIC
    except OSError:
        return False


def list_sheets(path: Path | str) -> list[SheetInfo]:
    """Sheets (and, for .xlsx, named Tables) in a file. CSV files have none."""
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".csv":
        return []
    try:
        if ext in _XLSX_EXTS:
            return _xlsx_sheets(path)
        if ext == ".xls":
            return _xls_sheets(path)
        if ext == ".ods":
            return _ods_sheets(path)
    except LoaderError:
        raise
    except Exception as exc:  # reader libraries raise many different types
        raise _wrap(path, exc) from exc
    raise LoaderError(f"Unsupported file type: {ext or path.name}")


def load_table(path: Path | str, sheet: str | None = None, table: str | None = None) -> pd.DataFrame:
    """Read a whole table into a DataFrame.

    ``sheet`` picks a sheet (first sheet when omitted). ``table`` picks a named
    Excel Table on that sheet (.xlsx only).
    """
    path = Path(path)
    ext = path.suffix.lower()
    try:
        if ext == ".csv":
            df = _read_csv(path)
        elif table is not None and ext in _XLSX_EXTS:
            df = _read_xlsx_table(path, sheet, table)
        elif ext in _XLSX_EXTS:
            df = pd.read_excel(path, sheet_name=sheet or 0, engine="openpyxl")
        elif ext == ".xls":
            df = pd.read_excel(path, sheet_name=sheet or 0, engine="xlrd")
        elif ext == ".ods":
            df = pd.read_excel(path, sheet_name=sheet or 0, engine="odf")
        else:
            raise LoaderError(f"Unsupported file type: {ext or path.name}")
    except LoaderError:
        raise
    except Exception as exc:  # reader libraries raise many different types
        raise _wrap(path, exc) from exc
    return clean_columns(df)


def _xlsx_sheets(path: Path) -> list[SheetInfo]:
    from openpyxl import load_workbook

    # Not read-only mode: named Tables are only available in normal mode.
    workbook = load_workbook(path, data_only=True)
    try:
        sheets = []
        for ws in workbook.worksheets:
            is_blank = ws.max_row == 1 and ws.max_column == 1 and ws["A1"].value is None
            rows, cols = (0, 0) if is_blank else (ws.max_row, ws.max_column)
            sheets.append(SheetInfo(ws.title, rows, cols, tuple(ws.tables.keys())))
        return sheets
    finally:
        workbook.close()


def _xls_sheets(path: Path) -> list[SheetInfo]:
    import xlrd

    book = xlrd.open_workbook(path, on_demand=True)
    try:
        return [SheetInfo(s.name, s.nrows, s.ncols) for s in (book.sheet_by_index(i) for i in range(book.nsheets))]
    finally:
        book.release_resources()


def _ods_sheets(path: Path) -> list[SheetInfo]:
    # Sizes are not available without reading each sheet, so only names are listed.
    with pd.ExcelFile(path, engine="odf") as workbook:
        return [SheetInfo(name) for name in workbook.sheet_names]


def _read_xlsx_table(path: Path, sheet: str | None, table: str) -> pd.DataFrame:
    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=True)
    try:
        ws = workbook[sheet] if sheet else workbook.worksheets[0]
        if table not in ws.tables:
            raise LoaderError(f"Table '{table}' not found on sheet '{ws.title}'")
        rows = [[cell.value for cell in row] for row in ws[ws.tables[table].ref]]
    finally:
        workbook.close()
    if not rows:
        return pd.DataFrame()
    return pd.DataFrame(rows[1:], columns=rows[0]).infer_objects()
